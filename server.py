from flask import Flask, jsonify, request, send_from_directory, session, redirect, url_for, send_file
import json, os, secrets, io, zipfile, shutil, re, csv
import pdfplumber
import openpyxl
import docx as python_docx
from datetime import datetime, timedelta
from werkzeug.security import generate_password_hash, check_password_hash

app = Flask(__name__, static_folder="static", template_folder="templates")
app.secret_key = os.environ.get("SECRET_KEY") or secrets.token_hex(32)
app.permanent_session_lifetime = timedelta(minutes=10)
app.config['SESSION_REFRESH_EACH_REQUEST'] = True  # sliding expiry: each request resets the 10-min clock

# ── LEGACY BOOTSTRAP CREDENTIALS ──────────────────────
# Used only once, to create the first admin account in users.json
# if that file doesn't exist yet. After that, all logins go
# through users.json — change your password from the Users tab.
BOOTSTRAP_USERNAME = "Sky Fr8"
BOOTSTRAP_PASSWORD = "Skyfr8@2022"

DATA_FILE    = os.path.join(os.path.dirname(__file__), "data", "shipments.json")
FLIGHTS_FILE = os.path.join(os.path.dirname(__file__), "data", "flights.json")
BIN_FILE     = os.path.join(os.path.dirname(__file__), "data", "bin.json")
USERS_FILE   = os.path.join(os.path.dirname(__file__), "data", "users.json")
ACTLOG_FILE  = os.path.join(os.path.dirname(__file__), "data", "activity_log.json")
ACTLOG_MAX_ENTRIES = 5000

BACKUP_FILES = {
    "shipments.json": DATA_FILE,
    "flights.json": FLIGHTS_FILE,
    "bin.json": BIN_FILE,
    "users.json": USERS_FILE,
    "activity_log.json": ACTLOG_FILE,
}
PRE_RESTORE_DIR = os.path.join(os.path.dirname(__file__), "data", "pre_restore_backups")

def load_json(path):
    if not os.path.exists(path):
        return []
    with open(path, "r") as f:
        return json.load(f)

def save_json(path, data):
    with open(path, "w") as f:
        json.dump(data, f, indent=2)

def is_logged_in():
    return session.get('logged_in') is True

def is_admin():
    return session.get('logged_in') is True and session.get('role') == 'admin'

def load_shipments(): return load_json(DATA_FILE)
def save_shipments(d): save_json(DATA_FILE, d)
def load_flights(): return load_json(FLIGHTS_FILE)
def save_flights(d): save_json(FLIGHTS_FILE, d)
def load_bin(): return load_json(BIN_FILE)
def save_bin(d): save_json(BIN_FILE, d)

def load_users(): return load_json(USERS_FILE)
def save_users(d): save_json(USERS_FILE, d)

def load_activity_log(): return load_json(ACTLOG_FILE)
def save_activity_log(d): save_json(ACTLOG_FILE, d)

def log_activity(action, detail=""):
    """Append an entry to the activity log. Uses the logged-in
    session's username, or 'system' if there isn't one."""
    entries = load_activity_log()
    entries.append({
        "timestamp": datetime.now().isoformat(),
        "username": session.get("username", "system"),
        "action": action,
        "detail": detail
    })
    if len(entries) > ACTLOG_MAX_ENTRIES:
        entries = entries[-ACTLOG_MAX_ENTRIES:]
    save_activity_log(entries)

def ensure_users_bootstrapped():
    """If users.json doesn't exist yet, create it with a single admin
    account using the old hardcoded credentials, so existing logins
    keep working after this update."""
    if os.path.exists(USERS_FILE):
        return
    bootstrap_admin = {
        "username": BOOTSTRAP_USERNAME,
        "password_hash": generate_password_hash(BOOTSTRAP_PASSWORD),
        "role": "admin",
        "active": True,
        "created": datetime.now().isoformat()
    }
    save_users([bootstrap_admin])

def find_user(username):
    return next((u for u in load_users() if u["username"].lower() == username.lower()), None)

def public_user(u):
    return {"username": u["username"], "role": u["role"], "active": u.get("active", True), "created": u.get("created")}

ensure_users_bootstrapped()

# ── MANIFEST PDF PARSER ───────────────────────────────
# Tuned to the Magaya Cargo System manifest layout: a repeating per-page
# header (agent/carrier/flight/tail/date) followed by a column-positioned
# list of shipment rows (no real table gridlines in the PDF, so columns
# are detected by x-position). This is a best-effort extraction — the
# frontend always shows results for staff review/edit before anything
# is saved.
_COL_SHIP=(75,225); _COL_CONS=(225,415); _COL_PCS=(415,485); _COL_WT=(485,550); _COL_GOODS=(550,670)
_TOTAL_ROW_CUTOFF = 545  # the page-bottom "TOTAL" summary is a fixed graphic, not real text

_AWB_RE    = re.compile(r'^\d{3}-\d{7,9}$')
_FLIGHT_RE = re.compile(r'^[A-Z]{2,4}\d{3,5}[A-Z]{0,6}(-[A-Z]{2,6})?$')
_TAIL_RE   = re.compile(r'^N\d{2,5}[A-Z]{0,3}$')
_DATE_RE   = re.compile(r'^[A-Za-z]{3}/\d{2}/\d{4}$')

def _col_of(x0):
    for lo, hi, name in [(*_COL_SHIP,'ship'), (*_COL_CONS,'cons'), (*_COL_PCS,'pcs'), (*_COL_WT,'wt'), (*_COL_GOODS,'goods')]:
        if lo <= x0 < hi: return name
    return None

def _group_lines(words, tol=2.5):
    words = sorted(words, key=lambda w: (w['top'], w['x0']))
    lines = []
    for w in words:
        if lines and abs(lines[-1][0]['top'] - w['top']) < tol:
            lines[-1].append(w)
        else:
            lines.append([w])
    return lines

def _parse_pieces(lines):
    for ln in lines:
        m = re.match(r'^(\d+)\s*(PCS|BOX|CTN|PLT)?', ln, re.I)
        if m: return int(m.group(1))
    return None

def _parse_weight_kg(lines):
    for ln in lines:
        m = re.match(r'^([\d,]+\.?\d*)\s*Kg', ln, re.I)
        if m: return float(m.group(1).replace(',', ''))
    return None

def _parse_phone_and_address(lines):
    phone = None
    addr_lines = []
    for ln in lines:
        m = re.match(r'^(Tel:)?\s*([\d\s\-+()]{7,})$', ln)
        if m and sum(c.isdigit() for c in ln) >= 7:
            phone = re.sub(r'[^\d+]', '', m.group(2))
        else:
            addr_lines.append(ln)
    return phone, addr_lines

def _guess_country(addr_lines):
    text = ' '.join(addr_lines).upper()
    if 'JAMAICA' in text: return 'Jamaica'
    if 'UNITED STATES' in text or re.search(r'\bFL\b', text): return 'United States'
    return ''

def _extract_manifest_page(page):
    words = page.extract_words()
    words = [w for w in words if w['top'] < _TOTAL_ROW_CUTOFF]

    flight_raw = next((w['text'] for w in words if _FLIGHT_RE.match(w['text']) and w['top'] < 200), None)
    tail       = next((w['text'] for w in words if _TAIL_RE.match(w['text']) and w['top'] < 200), None)
    dates      = [w['text'] for w in words if _DATE_RE.match(w['text']) and w['top'] < 200]
    eta_raw    = dates[0] if dates else None

    agent_words = sorted([w for w in words if w['x0'] < 160 and w['top'] < 130], key=lambda w: (w['top'], w['x0']))
    agent_name  = ' '.join(w['text'] for w in agent_words)

    airline_words = sorted([w for w in words if w['x0'] < 300 and w['top'] < 40], key=lambda w: w['x0'])
    airline = re.sub(r'\s*\([A-Z]{3}\)\s*$', '', ' '.join(w['text'] for w in airline_words)).strip()

    awb_words = sorted([w for w in words if _AWB_RE.match(w['text']) and w['x0'] < 80], key=lambda w: w['top'])

    blocks = []
    for i, aw in enumerate(awb_words):
        top_start = aw['top'] - 1
        top_end = awb_words[i+1]['top'] - 1 if i+1 < len(awb_words) else _TOTAL_ROW_CUTOFF
        blk_words = [w for w in words if top_start <= w['top'] < top_end]
        blocks.append((aw['text'], blk_words))

    eta_iso = None
    if eta_raw:
        try:
            eta_iso = datetime.strptime(eta_raw, '%b/%d/%Y').strftime('%Y-%m-%d')
        except ValueError:
            pass

    flight = None
    if flight_raw:
        m = re.match(r'^[A-Z]{2,4}\d{3,5}', flight_raw)
        flight = m.group(0) if m else flight_raw

    shipments = []
    for awb, blk_words in blocks:
        cols = {'ship': [], 'cons': [], 'pcs': [], 'wt': [], 'goods': []}
        for w in blk_words:
            c = _col_of(w['x0'])
            if c: cols[c].append(w)

        def col_lines(name):
            return [' '.join(w['text'] for w in ln) for ln in _group_lines(cols[name])]

        ship_lines, cons_lines = col_lines('ship'), col_lines('cons')
        pcs_lines, wt_lines, goods_lines = col_lines('pcs'), col_lines('wt'), col_lines('goods')

        sphone, ship_addr = _parse_phone_and_address(ship_lines)
        cphone, cons_addr = _parse_phone_and_address(cons_lines)

        shipments.append({
            "awb": awb,
            "is_skyfr8": "SKYFR8" in agent_name.upper(),
            "agent_name": agent_name,
            "flight": flight or "",
            "airline": airline or "",
            "registration": tail or "",
            "eta": eta_iso or "",
            "sname": ship_addr[0] if ship_addr else "",
            "s_street": ', '.join(ship_addr[1:]) if len(ship_addr) > 1 else "",
            "sphone": sphone or "",
            "s_country": _guess_country(ship_addr),
            "cname": cons_addr[0] if cons_addr else "",
            "c_street": ', '.join(cons_addr[1:]) if len(cons_addr) > 1 else "",
            "cphone": cphone or "",
            "c_country": _guess_country(cons_addr),
            "pcs_manifested": _parse_pieces(pcs_lines),
            "gw": _parse_weight_kg(wt_lines),
            "desc": ' '.join(goods_lines) if goods_lines else "COURIER GOODS",
        })
    return shipments

def parse_manifest_pdf(file_bytes):
    results = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            results.extend(_extract_manifest_page(page))
    return results

@app.route("/")
def index():
    if not is_logged_in():
        return redirect('/login')
    return send_from_directory("templates", "index.html")

@app.route("/login", methods=["GET"])
def login_page():
    return send_from_directory("templates", "login.html")

@app.route("/login", methods=["POST"])
def login():
    data = request.get_json()
    username = (data.get("username") or "").strip()
    password = data.get("password") or ""
    user = find_user(username)
    if not user or not check_password_hash(user["password_hash"], password):
        return jsonify({"error": "Invalid username or password."}), 401
    if not user.get("active", True):
        return jsonify({"error": "This account has been deactivated."}), 403
    session.permanent      = True
    session['logged_in'] = True
    session['username']  = user["username"]
    session['role']      = user["role"]
    log_activity("LOGIN")
    return jsonify({"success": True})

@app.route("/logout")
def logout():
    if is_logged_in():
        log_activity("LOGOUT")
    session.clear()
    return redirect('/login')

@app.route("/api/session", methods=["GET"])
def get_session():
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    return jsonify({"username": session.get("username"), "role": session.get("role")})

@app.route("/track")
def track():
    return send_from_directory("templates", "track.html")

# ── PUBLIC TRACKING (no login required) ───────────────
# Exposes only what the customer-facing /track page actually displays —
# no declared value, warehouse bay, internal remarks, or staff notification
# tracking fields, and no soft-deleted (recycle bin) shipments.
_TRACK_PUBLIC_FIELDS = [
    "awb","flight","airline","pol","pod","eta","asycuda","desc","type",
    "gw","pcs_manifested","status","pcs_delivered","date_delivered","customs",
    "sname","sphone","s_street","s_city","s_parish","s_zip","s_country",
    "cname","cphone","c_street","c_city","c_parish","c_zip","c_country",
    "hawbs",
]

@app.route("/api/track/shipments", methods=["GET"])
def track_shipments_public():
    shipments = load_shipments()
    public = []
    for s in shipments:
        row = {k: s.get(k, "") for k in _TRACK_PUBLIC_FIELDS}
        row["hawbs"] = [
            {k: h.get(k, "") for k in ("hawb","consignee","commodity","pcs")}
            for h in s.get("hawbs", [])
        ]
        public.append(row)
    return jsonify(public)

# ── SHIPMENTS ──────────────────────────────────────────
@app.route("/api/shipments", methods=["GET"])
def get_shipments():
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    return jsonify(load_shipments())

@app.route("/api/shipments", methods=["POST"])
def add_shipment():
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    shipments = load_shipments()
    new = request.get_json()
    awb = new.get("awb", "").strip()
    if not awb:
        return jsonify({"error": "AWB number is required."}), 400
    if any(s["awb"].lower() == awb.lower() for s in shipments):
        return jsonify({"error": "AWB already exists."}), 409
    flight = new.get("flight", "")
    date   = new.get("eta", "")
    match  = next((f for f in load_flights() if f["flight"].lower()==flight.lower() and f["date"]==date), None)
    if match and match.get("asycuda") and not new.get("asycuda"):
        new["asycuda"] = match["asycuda"]
    shipments.append(new)
    save_shipments(shipments)
    log_activity("ADD_SHIPMENT", f"AWB {awb}")
    return jsonify({"success": True, "awb": awb}), 201

@app.route("/api/shipments/<awb>", methods=["PUT"])
def update_shipment(awb):
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    shipments = load_shipments()
    data = request.get_json()
    for i, s in enumerate(shipments):
        if s["awb"].lower() == awb.lower():
            shipments[i].update(data)
            save_shipments(shipments)
            log_activity("EDIT_SHIPMENT", f"AWB {awb}")
            return jsonify({"success": True})
    return jsonify({"error": "AWB not found."}), 404

@app.route("/api/shipments/<awb>/bin", methods=["POST"])
def move_to_bin(awb):
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    shipments = load_shipments()
    record = next((s for s in shipments if s["awb"].lower()==awb.lower()), None)
    if not record:
        return jsonify({"error": "AWB not found."}), 404
    bin_items = load_bin()
    record["deleted_at"] = datetime.now().isoformat()
    bin_items.append(record)
    save_bin(bin_items)
    save_shipments([s for s in shipments if s["awb"].lower()!=awb.lower()])
    log_activity("DELETE_SHIPMENT", f"AWB {awb}")
    return jsonify({"success": True})

@app.route("/api/shipments/bulk-asycuda", methods=["POST"])
def bulk_asycuda():
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    data      = request.get_json()
    flight    = data.get("flight", "").strip()
    date      = data.get("date", "").strip()
    asycuda   = data.get("asycuda", "").strip()
    shipments = load_shipments()
    updated   = 0
    for s in shipments:
        if s.get("flight","").lower()==flight.lower() and s.get("eta","")==date:
            s["asycuda"] = asycuda
            updated += 1
    save_shipments(shipments)
    log_activity("BULK_ASYCUDA", f"Flight {flight} {date}: {updated} AWB(s) → {asycuda}")
    return jsonify({"success": True, "updated": updated})

# ── RECYCLE BIN ──────────────────────────────────────
@app.route("/api/bin", methods=["GET"])
def get_bin():
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    bin_items = load_bin()
    cutoff = datetime.now() - timedelta(days=30)
    bin_items = [b for b in bin_items if datetime.fromisoformat(b.get("deleted_at","2000-01-01")) > cutoff]
    save_bin(bin_items)
    return jsonify(bin_items)

@app.route("/api/bin/<awb>/restore", methods=["POST"])
def restore_from_bin(awb):
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    bin_items = load_bin()
    record = next((b for b in bin_items if b["awb"].lower()==awb.lower()), None)
    if not record:
        return jsonify({"error": "Record not found in bin."}), 404
    record.pop("deleted_at", None)
    shipments = load_shipments()
    shipments.append(record)
    save_shipments(shipments)
    save_bin([b for b in bin_items if b["awb"].lower()!=awb.lower()])
    log_activity("RESTORE_SHIPMENT", f"AWB {awb}")
    return jsonify({"success": True})

@app.route("/api/bin/<awb>", methods=["DELETE"])
def permanent_delete(awb):
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    bin_items = load_bin()
    new_bin = [b for b in bin_items if b["awb"].lower()!=awb.lower()]
    if len(new_bin)==len(bin_items):
        return jsonify({"error": "Record not found in bin."}), 404
    save_bin(new_bin)
    log_activity("PERM_DELETE", f"AWB {awb}")
    return jsonify({"success": True})

@app.route("/api/bin", methods=["DELETE"])
def empty_bin():
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    save_bin([])
    log_activity("EMPTY_BIN")
    return jsonify({"success": True})

# ── FLIGHTS ───────────────────────────────────────────
@app.route("/api/flights", methods=["GET"])
def get_flights():
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    return jsonify(load_flights())

@app.route("/api/flights", methods=["POST"])
def save_flight():
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    flights = load_flights()
    new = request.get_json()
    flight = new.get("flight","").strip()
    date   = new.get("date","").strip()
    if not flight or not date:
        return jsonify({"error": "Flight number and date are required."}), 400
    for i, f in enumerate(flights):
        if f["flight"].lower()==flight.lower() and f["date"]==date:
            flights[i].update(new)
            save_flights(flights)
            log_activity("EDIT_FLIGHT", f"{flight} ({date})")
            return jsonify({"success": True, "updated": True})
    flights.append(new)
    save_flights(flights)
    log_activity("ADD_FLIGHT", f"{flight} ({date})")
    return jsonify({"success": True, "updated": False}), 201

@app.route("/api/flights/<flight>/<date>", methods=["DELETE"])
def delete_flight(flight, date):
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    flights = load_flights()
    new_list = [f for f in flights if not (f["flight"].lower()==flight.lower() and f["date"]==date)]
    if len(new_list)==len(flights):
        return jsonify({"error": "Flight not found."}), 404
    save_flights(new_list)
    log_activity("DELETE_FLIGHT", f"{flight} ({date})")
    return jsonify({"success": True})

# ── USERS (admin only) ────────────────────────────────
@app.route("/api/users", methods=["GET"])
def get_users():
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    if not is_admin(): return jsonify({"error": "Admin access required."}), 403
    return jsonify([public_user(u) for u in load_users()])

@app.route("/api/users", methods=["POST"])
def create_user():
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    if not is_admin(): return jsonify({"error": "Admin access required."}), 403
    data     = request.get_json()
    username = (data.get("username") or "").strip()
    password = data.get("password") or ""
    role     = data.get("role") or "staff"
    if not username or not password:
        return jsonify({"error": "Username and password are required."}), 400
    if role not in ("admin", "staff"):
        return jsonify({"error": "Role must be 'admin' or 'staff'."}), 400
    if find_user(username):
        return jsonify({"error": "A user with that username already exists."}), 409
    users = load_users()
    users.append({
        "username": username,
        "password_hash": generate_password_hash(password),
        "role": role,
        "active": True,
        "created": datetime.now().isoformat()
    })
    save_users(users)
    log_activity("CREATE_USER", f"{username} ({role})")
    return jsonify({"success": True}), 201

@app.route("/api/users/<username>", methods=["PUT"])
def update_user(username):
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    if not is_admin(): return jsonify({"error": "Admin access required."}), 403
    data  = request.get_json()
    users = load_users()
    idx   = next((i for i,u in enumerate(users) if u["username"].lower()==username.lower()), None)
    if idx is None:
        return jsonify({"error": "User not found."}), 404

    target = users[idx]
    new_role   = data.get("role", target["role"])
    new_active = data.get("active", target.get("active", True))

    # Don't allow removing the last active admin's admin rights or access
    admins = [u for u in users if u["role"]=="admin" and u.get("active",True)]
    is_only_admin = target["role"]=="admin" and target.get("active",True) and len(admins)<=1
    if is_only_admin and (new_role != "admin" or not new_active):
        return jsonify({"error": "Can't remove the last active admin."}), 400

    target["role"]   = new_role
    target["active"] = new_active
    if data.get("password"):
        target["password_hash"] = generate_password_hash(data["password"])
    save_users(users)
    log_activity("EDIT_USER", f"{username} → role={new_role}, active={new_active}")
    return jsonify({"success": True})

@app.route("/api/users/<username>", methods=["DELETE"])
def delete_user(username):
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    if not is_admin(): return jsonify({"error": "Admin access required."}), 403
    if username.lower() == (session.get("username") or "").lower():
        return jsonify({"error": "You can't delete your own account while logged in."}), 400
    users  = load_users()
    target = next((u for u in users if u["username"].lower()==username.lower()), None)
    if not target:
        return jsonify({"error": "User not found."}), 404
    admins = [u for u in users if u["role"]=="admin" and u.get("active",True)]
    if target["role"]=="admin" and target.get("active",True) and len(admins)<=1:
        return jsonify({"error": "Can't delete the last active admin."}), 400
    save_users([u for u in users if u["username"].lower()!=username.lower()])
    log_activity("DELETE_USER", username)
    return jsonify({"success": True})

# ── ACTIVITY LOG (admin only) ─────────────────────────
@app.route("/api/activity-log", methods=["GET"])
def get_activity_log():
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    if not is_admin(): return jsonify({"error": "Admin access required."}), 403
    limit = request.args.get("limit", default=500, type=int)
    entries = load_activity_log()
    entries = list(reversed(entries))  # newest first
    if limit:
        entries = entries[:limit]
    return jsonify(entries)

# ── SPREADSHEET / WORD MANIFEST IMPORT ────────────────
# These formats vary from one sender to the next, so instead of guessing at a
# layout, the server returns the raw grid of cells and the browser lets staff
# map columns to fields before anything is reviewed or saved.
_TABLE_ROW_CAP = 5000

def _cell_to_str(v):
    if v is None: return ""
    if hasattr(v, "isoformat"):
        return v.strftime("%Y-%m-%d") if isinstance(v, datetime) else v.isoformat()
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip()

def parse_table_file(filename, data):
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text))]
    elif ext in ("xlsx", "xlsm"):
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        ws = wb.active
        rows = [[_cell_to_str(c) for c in r] for r in ws.iter_rows(values_only=True)]
    elif ext == "docx":
        d = python_docx.Document(io.BytesIO(data))
        if not d.tables:
            raise ValueError("No table was found in this Word document. The manifest needs to be laid out as a table.")
        biggest = max(d.tables, key=lambda t: len(t.rows))
        rows = []
        for r in biggest.rows:
            seen, cells = set(), []
            for c in r.cells:  # merged cells repeat; keep each distinct cell once
                if id(c._tc) in seen: continue
                seen.add(id(c._tc))
                cells.append(c.text.strip())
            rows.append(cells)
    else:
        raise ValueError("Unsupported file type. Use .xlsx, .csv or .docx.")

    rows = [r for r in rows if any(c for c in r)][:_TABLE_ROW_CAP]
    if not rows:
        raise ValueError("The file appears to be empty.")
    width = max(len(r) for r in rows)
    return [r + [""] * (width - len(r)) for r in rows]

@app.route("/api/import/table-preview", methods=["POST"])
def import_table_preview():
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    upload = request.files.get("file")
    if not upload or not upload.filename:
        return jsonify({"error": "No file uploaded."}), 400
    try:
        grid = parse_table_file(upload.filename, upload.read())
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        return jsonify({"error": f"Couldn't read that file: {e}"}), 400
    log_activity("PREVIEW_TABLE_IMPORT", f"{upload.filename}: {len(grid)} row(s) read")
    return jsonify({"grid": grid})

# ── MANIFEST PDF IMPORT ───────────────────────────────
@app.route("/api/import/manifest-preview", methods=["POST"])
def import_manifest_preview():
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401

    upload = request.files.get("file")
    if not upload or not upload.filename:
        return jsonify({"error": "No PDF uploaded."}), 400
    if not upload.filename.lower().endswith(".pdf"):
        return jsonify({"error": "Please upload a PDF file."}), 400

    try:
        rows = parse_manifest_pdf(upload.read())
    except Exception as e:
        return jsonify({"error": f"Couldn't read that PDF: {e}"}), 400

    if not rows:
        return jsonify({"error": "No shipment rows were recognised in this PDF. The layout may not match what this importer expects."}), 400

    existing_awbs = {s["awb"].lower() for s in load_shipments()}
    for r in rows:
        r["duplicate"] = r["awb"].lower() in existing_awbs

    log_activity("PREVIEW_MANIFEST_IMPORT", f"{upload.filename}: {len(rows)} row(s) found")
    return jsonify({"rows": rows})

@app.route("/api/import/manifest-confirm", methods=["POST"])
def import_manifest_confirm():
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    data = request.get_json()
    rows = data.get("rows", [])
    if not rows:
        return jsonify({"error": "No rows to import."}), 400

    shipments = load_shipments()
    existing_awbs = {s["awb"].lower() for s in shipments}
    flights = load_flights()

    added, skipped = [], []
    for r in rows:
        awb = (r.get("awb") or "").strip()
        if not awb:
            skipped.append({"awb": "(blank)", "reason": "No AWB"})
            continue
        if awb.lower() in existing_awbs:
            skipped.append({"awb": awb, "reason": "AWB already exists"})
            continue

        record = {
            "awb": awb,
            "flight": r.get("flight", ""),
            "airline": r.get("airline", ""),
            "registration": r.get("registration", ""),
            "pol": r.get("pol", "MIA"),
            "pod": r.get("pod", "KIN"),
            "eta": r.get("eta", ""),
            "asycuda": "",
            "desc": r.get("desc", ""),
            "type": "General",
            "gw": r.get("gw") or 0,
            "cw": 0,
            "pcs_manifested": r.get("pcs_manifested") or 0,
            "pcs_received": None,
            "bay": "",
            "status": "Manifested",
            "pcs_delivered": None,
            "date_delivered": None,
            "sname": r.get("sname", ""),
            "sphone": r.get("sphone", ""),
            "s_street": r.get("s_street", ""),
            "s_city": r.get("s_city", ""),
            "s_parish": r.get("s_parish", ""),
            "s_zip": r.get("s_zip", ""),
            "s_country": r.get("s_country", ""),
            "cname": r.get("cname", ""),
            "cphone": r.get("cphone", ""),
            "c_street": r.get("c_street", ""),
            "c_city": r.get("c_city", ""),
            "c_parish": r.get("c_parish", ""),
            "c_zip": r.get("c_zip", ""),
            "c_country": r.get("c_country", ""),
            "customs": "Not submitted",
            "val": 0,
            "bays": [],
            "hawbs": [],
        }
        shipments.append(record)
        existing_awbs.add(awb.lower())
        added.append(awb)

        # upsert the flight record if we have enough info and it doesn't already exist
        fl, dt = record["flight"], record["eta"]
        if fl and dt and not any(f["flight"].lower()==fl.lower() and f["date"]==dt for f in flights):
            flights.append({
                "flight": fl, "date": dt, "airline": record["airline"],
                "registration": record["registration"], "pol": record["pol"], "pod": record["pod"]
            })

    save_shipments(shipments)
    save_flights(flights)
    log_activity("IMPORT_MANIFEST", f"Imported {len(added)} shipment(s), skipped {len(skipped)}")

    return jsonify({"success": True, "added": added, "skipped": skipped})

# ── BACKUP & RESTORE (admin only) ─────────────────────
@app.route("/api/backup", methods=["GET"])
def download_backup():
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    if not is_admin(): return jsonify({"error": "Admin access required."}), 403

    mem = io.BytesIO()
    with zipfile.ZipFile(mem, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, path in BACKUP_FILES.items():
            data = load_json(path)
            zf.writestr(name, json.dumps(data, indent=2))
    mem.seek(0)

    log_activity("DOWNLOAD_BACKUP")
    stamp = datetime.now().strftime("%Y-%m-%d_%H%M")
    return send_file(mem, mimetype="application/zip", as_attachment=True,
                      download_name=f"skyfr8_backup_{stamp}.zip")

@app.route("/api/backup/restore", methods=["POST"])
def restore_backup():
    if not is_logged_in(): return jsonify({"error": "Unauthorised"}), 401
    if not is_admin(): return jsonify({"error": "Admin access required."}), 403

    upload = request.files.get("file")
    if not upload or not upload.filename:
        return jsonify({"error": "No backup file uploaded."}), 400

    try:
        zf = zipfile.ZipFile(io.BytesIO(upload.read()))
    except zipfile.BadZipFile:
        return jsonify({"error": "That file isn't a valid backup zip."}), 400

    names_in_zip = set(zf.namelist())
    known = [n for n in BACKUP_FILES if n in names_in_zip]
    if not known:
        return jsonify({"error": "This zip doesn't contain any recognised backup files (shipments.json, flights.json, bin.json, users.json, activity_log.json)."}), 400

    # Validate every file we're about to restore is valid JSON before touching anything
    parsed = {}
    for name in known:
        try:
            parsed[name] = json.loads(zf.read(name).decode("utf-8"))
        except (json.JSONDecodeError, UnicodeDecodeError):
            return jsonify({"error": f"{name} in that zip isn't valid JSON \u2014 restore cancelled, nothing was changed."}), 400

    # Safety snapshot of current data before we overwrite anything
    os.makedirs(PRE_RESTORE_DIR, exist_ok=True)
    snapshot_dir = os.path.join(PRE_RESTORE_DIR, datetime.now().strftime("%Y-%m-%d_%H%M%S"))
    os.makedirs(snapshot_dir, exist_ok=True)
    for name, path in BACKUP_FILES.items():
        if os.path.exists(path):
            shutil.copy2(path, os.path.join(snapshot_dir, name))

    # Prune old snapshots, keep the most recent 10
    snapshots = sorted(os.listdir(PRE_RESTORE_DIR))
    for old in snapshots[:-10]:
        shutil.rmtree(os.path.join(PRE_RESTORE_DIR, old), ignore_errors=True)

    restored_users = "users.json" in known
    for name in known:
        save_json(BACKUP_FILES[name], parsed[name])

    log_activity("RESTORE_BACKUP", f"Restored: {', '.join(known)}")

    if restored_users:
        session.clear()  # user accounts changed underneath the current session — force re-login

    return jsonify({"success": True, "restored": known, "logged_out": restored_users})

if __name__ == "__main__":
    ensure_users_bootstrapped()
    print("\n" + "="*52)
    print("  AirCargo WMS is running!")
    print("="*52)
    print("  Staff portal:    http://localhost:5000")
    print("  Customer track:  http://localhost:5000/track")
    print("")
    print("  To access from other devices on the network,")
    print("  replace 'localhost' with this computer's IP.")
    print("  (Run 'ipconfig' in Command Prompt to find it)")
    print("="*52 + "\n")
    app.run(host="0.0.0.0", port=5000, debug=False)
